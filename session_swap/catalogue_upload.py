"""Read the uploaded workbook without copying email or student data into JSON."""
import json
import re
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

from .importer import HEADERS, SIMPLE_HEADERS, simple_session, validate
from .service import Problem

INVITE_HEADERS = ['company name', 'event type', 'event date', 'event time', 'venue']
SUBJECT = re.compile(r'\bUpcoming\s+Event\s*\|\s*PGP\s+Co\s+2027\s*\|\s*Pre\s+Placement\s+Talk\s*:', re.I)


def heading(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().lower()


def event_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = re.sub(r'\s+', ' ', str(value or '').replace(',', ' ')).strip()
    for fmt in ('%Y-%m-%d', '%d %B %Y', '%d %b %Y', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise Problem('Use an Excel date, YYYY-MM-DD, DD/MM/YYYY or a date such as 27 October, 2026.')


def read_uploaded_workbook(path):
    from openpyxl import load_workbook
    path = Path(path)
    if path.suffix.lower() != '.xlsx' or path.stat().st_size > 10_000_000:
        raise Problem('Upload an .xlsx workbook smaller than 10 MB.')
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 30_000_000:
            raise Problem('Workbook exceeds the 30 MB uncompressed limit.')
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        supported = []
        for sheet in wb:
            if sheet.max_column and sheet.max_column > 50:
                continue
            headers = [heading(v) for v in next(sheet.iter_rows(values_only=True), [])]
            required = next((h for h in (INVITE_HEADERS, HEADERS['Sessions'], SIMPLE_HEADERS) if all(k in headers for k in h)), None)
            if required:
                supported.append((sheet, headers, required))
        preferred = [entry for entry in supported if entry[0].title == 'Sessions']
        choices = preferred or supported
        if len(choices) != 1:
            raise Problem('Use one session sheet, preferably named Sessions, with the documented columns.')
        sheet, headers, required = choices[0]
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise Problem('Remove duplicate column headings.')
        sessions = {}
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        for number, values in enumerate(rows, 2):
            if number > 20001:
                raise Problem('Limit the workbook to 20,000 data rows.')
            if not any(v is not None and str(v).strip() for v in values):
                continue
            def cell(name):
                i = headers.index(name)
                value = values[i] if i < len(values) else None
                if isinstance(value, str) and value.startswith('='):
                    raise Problem(f'Row {number}: formulas are not supported in imported columns.')
                return value
            if 'subject' in headers and not SUBJECT.search(str(cell('subject') or '')):
                continue
            record = {k: cell(k) for k in required}
            try:
                if required == INVITE_HEADERS:
                    company, kind = [str(record[k] or '').strip() for k in ('company name', 'event type')]
                    if not company or not kind:
                        raise Problem('Company Name and Event Type are required.')
                    record = dict(title=f'{company} · {kind}', room=str(record['venue'] or '').strip(),
                                  date=event_date(record['event date']), time=str(record['event time'] or '').strip())
                    session = simple_session(record, number)
                elif required == SIMPLE_HEADERS:
                    record['date'] = event_date(record['date'])
                    session = simple_session({k: str(v or '').strip() for k, v in record.items()}, number)
                else:
                    from datetime import time
                    record['date'] = event_date(record['date'])
                    session = {k: v.isoformat(timespec='minutes') if isinstance(v, time) else str(v or '').strip() for k, v in record.items()}
                if 'session_id' in headers and required != HEADERS['Sessions']:
                    explicit = str(cell('session_id') or '').strip()
                    if explicit:
                        session['session_id'] = session['course_id'] = explicit
                session['instructor'] = ''
                previous = sessions.get(session['session_id'])
                if previous and previous != session:
                    raise Problem('Conflicting rows identify the same session. Keep one current row per event.')
                sessions[session['session_id']] = session
            except Problem as error:
                raise Problem(f'Row {number}: {error}') from error
        # Validate every row before any network request. Only public event fields
        # and internal identifiers survive; never export the other workbook columns.
        return {'sessions': validate({'sessions': list(sessions.values())})['sessions']}
    finally:
        wb.close()


if __name__ == '__main__':
    try:
        payload = read_uploaded_workbook(sys.argv[1])
        Path(sys.argv[2]).write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
        print(f'Workbook validated: {len(payload["sessions"])} sessions.')
    except Exception:
        # Action logs can be public. Do not echo workbook values or file contents.
        print('Workbook validation failed. Check the sheet, columns, dates, time ranges, formulas and conflicting duplicate rows against EXCEL_UPLOAD_SETUP.md.', file=sys.stderr)
        sys.exit(1)
