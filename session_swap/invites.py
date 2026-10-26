"""Parse saved Outlook invites into reviewable Excel session rows.

This module does not sign into Outlook or scrape the desktop app.
"""
import hashlib
import json
import re
import unicodedata
from datetime import datetime
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path

from .service import Problem


class PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.ignore = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.ignore += 1
        if tag in ('br', 'p', 'div', 'tr', 'li'):
            self.parts.append('\n')
        elif tag in ('td', 'th'):
            self.parts.append(' ')

    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self.ignore = max(0, self.ignore - 1)
        if tag in ('p', 'div', 'tr', 'li'):
            self.parts.append('\n')

    def handle_data(self, data):
        if not self.ignore:
            self.parts.append(data)


def parse_text(text, section='Main'):
    if not isinstance(text, str) or len(text) > 500000:
        raise Problem('Invite text must be shorter than 500,000 characters.')
    text = unicodedata.normalize('NFKC', text).replace('\r', '').replace('\u00a0', ' ')
    if re.search(r'<(?:html|body|div|p|table|br)\b', text, re.I):
        parser = PlainHTML()
        parser.feed(text)
        text = ''.join(parser.parts)
    labels = ['Company Name', 'Industry', 'CBD Name', 'Event Type', 'Event Date', 'Event Time', 'Duration', 'Mode', 'Venue', "Presenter's Designation", "Presenter’s Designation", "Presenter's Profile", "Presenter’s Profile"]
    pattern = r'(?:^|\n)\s*(' + '|'.join(re.escape(label) for label in labels) + r')\s*:\s*'
    matches = list(re.finditer(pattern, text, flags=re.I))
    fields = {}
    for i, match in enumerate(matches):
        value = text[match.end():matches[i+1].start() if i+1 < len(matches) else len(text)].strip()
        value = value.split('\n\n')[0].strip()
        fields[match.group(1).lower().replace('’', "'")] = value
    required = ['company name', 'event type', 'event date', 'event time']
    missing = [key for key in required if not fields.get(key)]
    if missing:
        raise Problem('Invite needs review: missing ' + ', '.join(missing) + '.')
    parsed_date = None
    for fmt in ('%d %B, %Y', '%d %B %Y', '%d %b, %Y', '%Y-%m-%d'):
        try:
            parsed_date = datetime.strptime(fields['event date'].strip().rstrip('.'), fmt).date()
            break
        except ValueError:
            pass
    if not parsed_date:
        raise Problem('Invite needs review: unrecognised event date.')
    match = re.fullmatch(r'\s*(\d{1,2}:\d{2})\s*[-–—]\s*(\d{1,2}:\d{2})\s*', fields['event time'])
    if not match:
        raise Problem('Invite needs review: event time must contain a start and end in 24-hour HH:MM format.')
    try:
        start = datetime.strptime(match[1], '%H:%M').time()
        end = datetime.strptime(match[2], '%H:%M').time()
    except ValueError:
        raise Problem('Invite needs review: invalid event time.')
    if end <= start:
        raise Problem('Invite needs review: end must follow start on the same day.')
    if not section.strip() or len(section) > 100:
        raise Problem('Provide a section label of 1–100 characters, or Main for a single event.')
    slug = re.sub(r'[^a-z0-9]+', '-', (fields['company name'] + '-' + fields['event type']).lower()).strip('-')[:80]
    course_id = f'{slug}-{parsed_date.isoformat()}'
    suffix = hashlib.sha256(section.strip().encode()).hexdigest()[:8]
    row = dict(session_id=f'{course_id}-{suffix}', course_id=course_id,
               title=fields['company name'] + ' · ' + fields['event type'], section=section.strip(),
               date=parsed_date.isoformat(), start_time=start.strftime('%H:%M'), end_time=end.strftime('%H:%M'),
               room=fields.get('venue', 'TBD'), instructor=fields.get("presenter's designation", ''))
    return dict(session=row, extracted=dict(company=fields['company name'], industry=fields.get('industry', ''),
                event_type=fields['event type'], mode=fields.get('mode', ''), venue=fields.get('venue', 'TBD'),
                presenter_designation=fields.get("presenter's designation", '')),
                review_required=True, notes=['Confirm the event, section and timezone before importing.',
                                            'Student names and PGIDs are entered in the app; they are not extracted from the email.'])


def read_invite(path, section='Main'):
    path = Path(path)
    if path.stat().st_size > 5_000_000:
        raise Problem('Invite must be smaller than 5 MB.')
    if path.suffix.lower() == '.eml':
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        body = message.get_body(preferencelist=('plain', 'html'))
        if not body:
            raise Problem('No email body found. Calendar-only attachments need separate review.')
        text = body.get_content()
    else:
        text = path.read_text(encoding='utf-8-sig')
    return parse_text(text, section)


def append_to_workbook(path, result):
    """Application runtime: update one reviewed event without losing other rows."""
    import os
    import tempfile
    from openpyxl import Workbook, load_workbook
    from .importer import HEADERS, read_excel, simple_session
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise Problem('Choose an .xlsx destination.')
    if path.exists():
        read_excel(path)
        workbook = load_workbook(path)
    else:
        workbook = Workbook()
        workbook.active.title = 'Sessions'
        workbook.active.append(HEADERS['Sessions'])
    sheet = workbook['Sessions']
    headers = [str(cell.value).strip().lower() if cell.value is not None else '' for cell in sheet[1]]
    simple = 'time' in headers and 'session_id' not in headers
    if result.get('source') and not simple:
        extra_headers = ['company', 'industry', 'event_type', 'mode', 'source_message_id', 'source_received_at', 'source_subject']
        existing = [cell.value for cell in sheet[1]]
        for heading in extra_headers:
            if heading not in existing:
                sheet.cell(1, sheet.max_column + 1, heading)
    headers = [str(cell.value).strip().lower() if cell.value is not None else '' for cell in sheet[1]]
    row = dict(result['session'])
    extracted = result.get('extracted', {})
    for heading in ('company', 'industry', 'event_type', 'mode'):
        if heading in extracted:
            row[heading] = extracted[heading]
    source = result.get('source', {})
    for key, heading in [('id', 'source_message_id'), ('received_at', 'source_received_at'), ('subject', 'source_subject')]:
        if key in source:
            row[heading] = source[key]
    if simple:
        row['time'] = f'{row["start_time"]} - {row["end_time"]}'
        normalized = simple_session(row, 2)
        current = read_excel(path)['sessions']
        # read_excel skips empty rows, so map back to the actual worksheet rows.
        occupied = [i for i in range(2, sheet.max_row+1) if any(
            cell.value is not None and str(cell.value).strip() for cell in sheet[i])]
        target = next((i for i, event in zip(occupied, current)
                       if event['session_id'] == normalized['session_id']), sheet.max_row+1)
    else:
        index = headers.index('session_id') + 1
        target = next((i for i in range(2, sheet.max_row+1) if sheet.cell(i,index).value == row['session_id']), sheet.max_row+1)
    for column, heading in enumerate(headers, 1):
        if heading in row:
            value = row[heading]
            if heading == 'date':
                value = datetime.strptime(value, '%Y-%m-%d').date()
            elif heading in ('start_time', 'end_time'):
                value = datetime.strptime(value, '%H:%M').time()
            cell = sheet.cell(target, column)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = 's'
            if heading == 'date':
                sheet.cell(target,column).number_format = 'dd mmm yyyy'
            elif heading in ('start_time','end_time'):
                sheet.cell(target,column).number_format = 'hh:mm'
    for table in sheet.tables.values():
        table.ref = f'A1:{sheet.cell(1,len(headers)).column_letter}{sheet.max_row}'
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = tempfile.NamedTemporaryFile(suffix='.xlsx', dir=path.parent, delete=False)
    temp.close()
    try:
        workbook.save(temp.name)
        read_excel(temp.name)
        os.replace(temp.name, path)
    finally:
        workbook.close()
        Path(temp.name).unlink(missing_ok=True)
    return row['session_id']
