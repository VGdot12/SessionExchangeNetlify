"""Validate complete Excel snapshots before making any database changes."""
import hashlib
import json
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from .db import set_meta, transaction
from .service import Problem, future, now, session

HEADERS = {
    'Sessions': ['session_id', 'course_id', 'title', 'section', 'date', 'start_time', 'end_time', 'room', 'instructor'],
    'Students': ['pgid', 'name'],
    'Assignments': ['pgid', 'session_id'],
}


def read_excel(path):
    from openpyxl import load_workbook
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise Problem('Use an .xlsx workbook. Export older .xls workbooks as .xlsx first.')
    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > 30_000_000:
            raise Problem('Workbook is too large. Limit uncompressed contents to 30 MB.')
    workbook = load_workbook(path, read_only=True, data_only=False)
    output = {}
    try:
        for sheet_name, required in HEADERS.items():
            if sheet_name not in workbook.sheetnames:
                if sheet_name == 'Sessions':
                    raise Problem(f'Missing sheet: {sheet_name}.')
                continue
            sheet = workbook[sheet_name]
            if sheet.max_column and sheet.max_column > 50:
                raise Problem(f'{sheet_name}: limit the sheet to 50 columns.')
            rows = sheet.iter_rows(values_only=True)
            headers = [str(v).strip().lower() if v is not None else '' for v in next(rows, [])]
            if len([h for h in headers if h]) != len(set(h for h in headers if h)):
                raise Problem(f'{sheet_name}: duplicate column headings.')
            if any(h not in headers for h in required):
                raise Problem(f'{sheet_name}: required columns are {", ".join(required)}.')
            items = []
            for row_number, values in enumerate(rows, 2):
                if row_number > 20002:
                    raise Problem(f'{sheet_name}: limit the sheet to 20,000 data rows.')
                if not any(v is not None and str(v).strip() for v in values):
                    continue
                record = {}
                for key in required:
                    value = values[headers.index(key)] if headers.index(key) < len(values) else None
                    if isinstance(value, str) and value.startswith('='):
                        raise Problem(f'{sheet_name} row {row_number}: formulas are not allowed in imported columns.')
                    if isinstance(value, datetime):
                        value = value.date().isoformat() if key == 'date' else value.time().isoformat(timespec='minutes')
                    elif isinstance(value, (date, time)):
                        value = value.isoformat()
                    elif isinstance(value, float) and value.is_integer():
                        value = str(int(value))
                    record[key] = '' if value is None else str(value).strip()
                items.append(record)
            output[sheet_name.lower()] = items
        return validate(output)
    finally:
        workbook.close()


def validate(payload):
    if not isinstance(payload, dict):
        raise Problem('Import must be a workbook snapshot object.')
    for key in ('sessions', 'students', 'assignments'):
        if key == 'sessions' and not payload.get(key):
            raise Problem('sessions: provide at least one event.')
        if not isinstance(payload.get(key, []), list) or len(payload.get(key, [])) > 20000:
            raise Problem(f'{key}: provide at most 20,000 rows.')
    normalized = {'sessions': [], 'students': [], 'assignments': []}
    for sheet_name, headings in HEADERS.items():
        key = sheet_name.lower()
        for number, row in enumerate(payload.get(key, []), 2):
            if not isinstance(row, dict):
                raise Problem(f'{sheet_name} row {number}: invalid record.')
            clean = {}
            for heading in headings:
                value = row.get(heading, '')
                if not isinstance(value, str) or len(value) > 200:
                    raise Problem(f'{sheet_name} row {number}: {heading} must be text of 200 characters or fewer.')
                clean[heading] = value.strip()
                if heading not in ('room', 'instructor') and not clean[heading]:
                    raise Problem(f'{sheet_name} row {number}: {heading} is required.')
            normalized[key].append(clean)
    sessions, students = {}, {}
    sections, course_titles = set(), {}
    for row in normalized['sessions']:
        sid = row['session_id']
        if sid in sessions:
            raise Problem(f'Duplicate session_id: {sid}.')
        try:
            row['date'] = date.fromisoformat(row['date']).isoformat()
            start, end = time.fromisoformat(row['start_time']), time.fromisoformat(row['end_time'])
            if start.tzinfo or end.tzinfo or end <= start or start.second or end.second or start.microsecond or end.microsecond:
                raise ValueError()
            row['start_time'], row['end_time'] = start.isoformat(timespec='minutes'), end.isoformat(timespec='minutes')
        except ValueError:
            raise Problem(f'{sid}: use YYYY-MM-DD dates and HH:MM times, with end after start on the same day.')
        section_key = (row['course_id'], row['section'])
        if section_key in sections:
            raise Problem(f'Duplicate section {row["section"]} for {row["course_id"]}. Use a separate course_id for each session topic or occurrence.')
        sections.add(section_key)
        if row['course_id'] in course_titles and course_titles[row['course_id']] != row['title']:
            raise Problem(f'{row["course_id"]}: every section must have the same title.')
        course_titles[row['course_id']] = row['title']
        sessions[sid] = row
    for row in normalized['students']:
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,40}', row['pgid']):
            raise Problem('PGIDs must be 1–40 letters, digits, underscores or hyphens. Store them as Excel text to preserve leading zeros.')
        if row['pgid'] in students:
            raise Problem(f'Duplicate PGID: {row["pgid"]}.')
        students[row['pgid']] = row
    seen = set()
    for row in normalized['assignments']:
        if row['pgid'] not in students or row['session_id'] not in sessions:
            raise Problem(f'Assignment {row["pgid"]}/{row["session_id"]} references a missing student or session.')
        key = (row['pgid'], sessions[row['session_id']]['course_id'])
        if key in seen:
            raise Problem(f'{row["pgid"]}: assign only one section per session topic.')
        seen.add(key)
    return normalized


def import_snapshot(conn, payload, source='Excel sync'):
    payload = validate(payload)
    digest = hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()
    sessions = {r['session_id']: r for r in payload['sessions']}
    with transaction(conn):
        # A stable session_id must always refer to the same topic.
        for row in payload['sessions']:
            old = conn.execute('SELECT course_id FROM sessions WHERE id=?', (row['session_id'],)).fetchone()
            if old and old['course_id'] != row['course_id']:
                raise Problem(f'{row["session_id"]}: course_id changed. Create a new session_id instead.')
        conn.execute('UPDATE sessions SET active=0')
        for r in payload['students']:
            conn.execute('INSERT INTO students(pgid,name) VALUES (?,?) ON CONFLICT(pgid) DO UPDATE SET name=excluded.name,active=1', (r['pgid'], r['name']))
        for r in payload['sessions']:
            conn.execute('INSERT INTO sessions(id,course_id,title,section,date,start_time,end_time,room,instructor) VALUES (?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET title=excluded.title,section=excluded.section,date=excluded.date,start_time=excluded.start_time,end_time=excluded.end_time,room=excluded.room,instructor=excluded.instructor,active=1', tuple(r[k] for k in HEADERS['Sessions']))
        for r in payload['assignments']:
            slot_id = r['session_id']
            # Preserve swaps when Excel's original assignment has not changed.
            conn.execute('INSERT INTO assignments VALUES (?,?,?,?) ON CONFLICT(pgid,slot_id) DO NOTHING', (r['pgid'], slot_id, r['session_id'], r['session_id']))
        cancelled = 0
        for r in conn.execute("SELECT * FROM requests WHERE status='OPEN'").fetchall():
            offer = session(conn, r['offered_id'])
            desired = session(conn, r['desired_id']) if r['desired_id'] else None
            any_alternative = any(s['id'] != offer['id'] and future(s) for s in conn.execute('SELECT * FROM sessions WHERE active=1'))
            if not offer['active'] or (desired and not desired['active']) or (not desired and not any_alternative):
                conn.execute("UPDATE requests SET status='CANCELLED',closed_at=?,close_reason='Event changed during Excel sync' WHERE id=?", (now(), r['id']))
                cancelled += 1
        result = dict(synced_at=now(), source=source, digest=digest, sessions=len(payload['sessions']), students=len(payload['students']), assignments=len(payload['assignments']), cancelled_requests=cancelled)
        set_meta(conn, 'sync', result)
    return result
