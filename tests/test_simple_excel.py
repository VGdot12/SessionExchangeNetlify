import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from session_swap.importer import read_excel
from session_swap.invites import append_to_workbook, parse_text
from session_swap.automation import import_messages
from session_swap.service import Problem
from test_app import INVITE


class SimpleExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / 'sessions.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def write(self, rows):
        wb = Workbook()
        wb.active.title = 'Sessions'
        wb.active.append(['title', 'room', 'time', 'date'])
        for row in rows:
            wb.active.append(row)
        wb.save(self.path)
        wb.close()

    def test_four_columns_and_stable_identity_across_room_end_and_row_changes(self):
        self.write([['Nagarro · Pre Placement Talk', 'TBD', '21:30 - 22:30', date(2026, 10, 27)]])
        first = read_excel(self.path)['sessions'][0]
        self.assertEqual((first['start_time'], first['end_time']), ('21:30', '22:30'))
        self.write([['Other event', 'LT 2', '9:00 – 10:00', '2026-10-28'],
                    ['Nagarro · Pre Placement Talk', 'LT 3', '21:30 - 23:00', '2026-10-27']])
        changed = read_excel(self.path)['sessions'][1]
        self.assertEqual(first['session_id'], changed['session_id'])
        self.assertEqual(changed['room'], 'LT 3')
        self.assertEqual(changed['instructor'], '')

    def test_invalid_ranges_duplicates_and_formulas_are_rejected(self):
        row = ['Event', 'LT 1', '18:00 - 19:00', '2026-10-28']
        for value in ['18:00', '25:00 - 26:00', '19:00 - 18:00', '=A1']:
            self.write([[row[0], row[1], value, row[3]]])
            with self.assertRaises(Problem):
                read_excel(self.path)
        self.write([row, row])
        with self.assertRaises(Problem):
            read_excel(self.path)

    def test_invite_updates_simple_workbook_without_adding_columns(self):
        self.write([['Nagarro · Pre Placement Talk', 'TBD', '21:30 - 22:30', '2026-10-27']])
        append_to_workbook(self.path, parse_text(INVITE.replace('Venue : TBD', 'Venue : LT 3')))
        self.assertEqual(len(read_excel(self.path)['sessions']), 1)
        self.assertEqual(read_excel(self.path)['sessions'][0]['room'], 'LT 3')
        wb = load_workbook(self.path, read_only=True)
        self.assertEqual(wb['Sessions'].max_column, 4)
        wb.close()

    def test_outlook_checkpoint_prevents_older_mail_regressing_simple_workbook(self):
        self.write([['Nagarro · Pre Placement Talk', 'TBD', '21:30 - 22:30', '2026-10-27']])
        checkpoint = Path(self.temp.name) / 'state.json'
        message = dict(id='new', sender='casevents@isb.edu', subject='Upcoming Event | Talk',
                       received_at='2026-10-26T12:00:00Z', body=INVITE.replace('Venue : TBD', 'Venue : LT 3'))
        import_messages([message], self.path, checkpoint, apply=True)
        message.update(id='old', received_at='2026-10-26T10:00:00Z', body=INVITE)
        self.assertEqual(import_messages([message], self.path, checkpoint, apply=True)['skipped'], 1)
        self.assertEqual(read_excel(self.path)['sessions'][0]['room'], 'LT 3')


if __name__ == '__main__':
    unittest.main()
