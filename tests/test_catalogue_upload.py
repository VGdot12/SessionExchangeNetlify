import tempfile
import unittest
from pathlib import Path
from datetime import date

from openpyxl import Workbook
from session_swap.catalogue_upload import read_uploaded_workbook
from session_swap.service import Problem

SUBJECT = 'Upcoming Event  | PGP Co 2027 | Pre Placement Talk : Session Scheduled'
HEADERS = ['Subject\u00a0', 'Company Name', 'Industry\u00a0', 'CBD Name', 'Event Type', 'Event Date', 'Event Time', 'Duration', 'Mode', 'Venue', 'Received On']


class CatalogueUploadTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / 'sessions.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def write(self, headers, rows, name='Table1'):
        wb = Workbook()
        wb.active.title = name
        wb.active.append(headers)
        for row in rows:
            wb.active.append(row)
        wb.save(self.path)
        wb.close()

    def invite(self, **changes):
        values = dict(zip(HEADERS, [SUBJECT, 'Nagarro', 'PRIVATE INDUSTRY', 'PRIVATE PERSON', 'Pre Placement Talk', '27 October, 2026', '21:30 - 22:30', '1 hr', 'Offline', 'LT 3', 'PRIVATE RECEIVED']))
        values.update(changes)
        return [values[h] for h in HEADERS]

    def test_actual_headings_named_date_and_private_columns_are_handled(self):
        self.write(HEADERS, [self.invite()])
        payload = read_uploaded_workbook(self.path)
        event = payload['sessions'][0]
        self.assertEqual(event['title'], 'Nagarro · Pre Placement Talk')
        self.assertEqual((event['room'], event['date'], event['start_time'], event['end_time']), ('LT 3', '2026-10-27', '21:30', '22:30'))
        self.assertNotIn('PRIVATE', str(payload))
        self.assertNotIn('Subject', str(payload))

    def test_excel_date_and_simplified_format_keep_stable_ids(self):
        self.write(HEADERS, [self.invite(**{'Event Date': date(2026, 10, 27)})])
        first = read_uploaded_workbook(self.path)['sessions'][0]
        self.write(['title', 'room', 'time', 'date'], [['Nagarro · Pre Placement Talk', 'LT 4', '21:30 – 23:00', '27/10/2026']], 'Sessions')
        second = read_uploaded_workbook(self.path)['sessions'][0]
        self.assertEqual(first['session_id'], second['session_id'])
        self.assertEqual(second['room'], 'LT 4')

    def test_subject_filter_ignores_other_cohorts_and_collapses_identical_rows(self):
        self.write(HEADERS, [self.invite(), self.invite(), self.invite(**{'Subject\u00a0': SUBJECT.replace('2027', '2028'), 'Company Name': 'Other'})])
        self.assertEqual(len(read_uploaded_workbook(self.path)['sessions']), 1)
        self.write(HEADERS, [self.invite(**{'Subject\u00a0': 'Other email'})])
        with self.assertRaises(Problem):
            read_uploaded_workbook(self.path)

    def test_invalid_rows_formulas_and_conflicting_duplicate_events_fail(self):
        for changes in [{'Event Time': '23:00 - 01:00'}, {'Event Date': '2026-02-30'}, {'Venue': '=A1'}, {'Company Name': ''}]:
            self.write(HEADERS, [self.invite(**changes)])
            with self.assertRaises(Problem):
                read_uploaded_workbook(self.path)
        self.write(HEADERS, [self.invite(), self.invite(**{'Venue': 'Other room'})])
        with self.assertRaises(Problem):
            read_uploaded_workbook(self.path)

    def test_original_schema_and_optional_stable_ids_are_supported(self):
        from session_swap.importer import HEADERS as ORIGINAL
        self.write(ORIGINAL['Sessions'], [['legacy', 'topic', 'Legacy', 'A', '2026-10-27', '20:00', '21:00', 'LT 1', 'PRIVATE PERSON']])
        event = read_uploaded_workbook(self.path)['sessions'][0]
        self.assertEqual((event['session_id'], event['course_id'], event['instructor']), ('legacy', 'topic', ''))
        self.write(['title', 'room', 'time', 'date', 'session_id'], [['Stable event', 'LT 1', '20:00 - 21:00', '2026-10-27', 'stable-1']])
        self.assertEqual(read_uploaded_workbook(self.path)['sessions'][0]['session_id'], 'stable-1')

    def test_ambiguous_sheets_fail_and_named_sessions_sheet_wins(self):
        self.write(['title', 'room', 'time', 'date'], [['One', 'LT 1', '20:00 - 21:00', '2026-10-27']])
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        wb.copy_worksheet(wb.active)
        wb.save(self.path)
        wb.close()
        with self.assertRaises(Problem):
            read_uploaded_workbook(self.path)
        wb = load_workbook(self.path)
        wb.active.title = 'Sessions'
        wb.save(self.path)
        wb.close()
        self.assertEqual(len(read_uploaded_workbook(self.path)['sessions']), 1)


if __name__ == '__main__':
    unittest.main()
